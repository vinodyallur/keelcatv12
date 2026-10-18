#!/usr/bin/env python3
"""
KeelCat — Techstars pitch deck generator.

Builds a clean, on-brand 16:9 investor deck (KeelCat-Pitch-Deck.pptx) using
python-pptx. Brand: dark surface + lime accent, matching the product UI.

Run:  python apisteward/docs/build_pitch_deck.py
Needs: pip install python-pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- Brand palette ---------------------------------------------------------
BG       = RGBColor(0x12, 0x14, 0x1A)   # deep near-black surface
CARD     = RGBColor(0x1B, 0x1E, 0x26)   # slightly lifted card
ACCENT   = RGBColor(0xA8, 0xFF, 0x53)   # KeelCat lime
INK      = RGBColor(0xF4, 0xF6, 0xF9)   # near-white text
MUTED    = RGBColor(0x9A, 0xA3, 0xB2)   # muted gray
FAINT    = RGBColor(0x66, 0x6C, 0x78)   # faint gray
ROSE     = RGBColor(0xFF, 0x6B, 0x7A)   # warning/problem
FONT     = "Segoe UI"

EMU_W, EMU_H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]


# ---- Helpers ---------------------------------------------------------------
def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    # send to back
    sp = bg._element
    sp.getparent().remove(sp)
    s.shapes._spTree.insert(2, sp)
    return s


def _set(tf_para, text, size, color, bold=False, italic=False):
    tf_para.text = text
    for r in tf_para.runs:
        r.font.size = Pt(size); r.font.color.rgb = color
        r.font.bold = bold; r.font.italic = italic; r.font.name = FONT


def textbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf


def rect(s, x, y, w, h, fill=CARD, line=None):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    # soften the rounded corner
    try:
        shp.adjustments[0] = 0.06
    except Exception:
        pass
    return shp


def eyebrow(s, text, x=Inches(0.7), y=Inches(0.55)):
    tf = textbox(s, x, y, Inches(11), Inches(0.4))
    _set(tf.paragraphs[0], text.upper(), 12, ACCENT, bold=True)
    tf.paragraphs[0].runs[0].font._rPr.set("spc", "200")


def title(s, text, x=Inches(0.7), y=Inches(0.95), size=34, w=Inches(11.9)):
    tf = textbox(s, x, y, w, Inches(1.2))
    _set(tf.paragraphs[0], text, size, INK, bold=True)
    return tf


def bullets(s, items, x=Inches(0.7), y=Inches(2.1), w=Inches(11.9), size=18, gap=10):
    tf = textbox(s, x, y, w, Inches(4.6))
    for i, (head, body) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        run = p.add_run(); run.text = "▸  "
        run.font.size = Pt(size); run.font.color.rgb = ACCENT; run.font.bold = True; run.font.name = FONT
        run = p.add_run(); run.text = head
        run.font.size = Pt(size); run.font.color.rgb = INK; run.font.bold = True; run.font.name = FONT
        if body:
            run = p.add_run(); run.text = "  —  " + body
            run.font.size = Pt(size); run.font.color.rgb = MUTED; run.font.name = FONT
    return tf


def footer(s, page):
    tf = textbox(s, Inches(0.7), Inches(7.02), Inches(12), Inches(0.35))
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = "KeelCat"; r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = ACCENT; r.font.name = FONT
    r = p.add_run(); r.text = "   ·   keelcat.in"; r.font.size = Pt(10); r.font.color.rgb = FAINT; r.font.name = FONT
    t2 = textbox(s, Inches(12.2), Inches(7.02), Inches(0.9), Inches(0.35))
    _set(t2.paragraphs[0], str(page), 10, FAINT); t2.paragraphs[0].alignment = PP_ALIGN.RIGHT


# ---- 1. Title -------------------------------------------------------------
s = slide()
rect(s, Inches(0.7), Inches(2.35), Inches(0.14), Inches(2.0), fill=ACCENT)
tf = textbox(s, Inches(1.05), Inches(2.2), Inches(11), Inches(1.2))
_set(tf.paragraphs[0], "KeelCat", 60, INK, bold=True)
tf = textbox(s, Inches(1.05), Inches(3.5), Inches(11.2), Inches(1.0))
_set(tf.paragraphs[0], "Your APIs change. Your code keeps up.", 26, ACCENT, bold=True)
tf = textbox(s, Inches(1.05), Inches(4.35), Inches(11), Inches(0.8))
_set(tf.paragraphs[0], "Dependabot for the APIs your code calls — it detects breaking API changes,\nfixes your code across every repo, and opens a pull request you review.", 16, MUTED)
tf = textbox(s, Inches(1.05), Inches(6.4), Inches(11.5), Inches(0.6))
p = tf.paragraphs[0]
r = p.add_run(); r.text = "Vinod & Srushti"; r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = INK; r.font.name = FONT
r = p.add_run(); r.text = "     ·     keelcat.in     ·     Techstars Application 2026"; r.font.size = Pt(14); r.font.color.rgb = FAINT; r.font.name = FONT

# ---- 2. Problem -----------------------------------------------------------
s = slide(); eyebrow(s, "The problem"); title(s, "APIs break your code silently — in production")
bullets(s, [
    ("Modern software runs on other people's APIs", "Stripe, Twilio, OpenAI, Anthropic, and dozens of internal services."),
    ("Every one of them changes", "a field is renamed, a method removed, an endpoint deprecated."),
    ("You find out when it breaks", "usually in production, for a customer, at the worst time."),
    ("The fix is a manual fire drill", "hunt every call site, across every repo and language, and rewrite by hand."),
])
# stat band
rect(s, Inches(0.7), Inches(5.55), Inches(11.9), Inches(1.15), fill=CARD)
for i,(num,label) in enumerate([("~40%","of integration failures come from\nunmanaged API changes"),
                                ("15–20 hrs","of emergency engineering\nper incident"),
                                ("~90 days","AI API deprecation cycle —\nand shrinking")]):
    tf = textbox(s, Inches(0.9)+Inches(3.95)*i, Inches(5.62), Inches(3.8), Inches(1.0), anchor=MSO_ANCHOR.MIDDLE)
    _set(tf.paragraphs[0], num, 24, ACCENT, bold=True)
    p = tf.add_paragraph(); _set(p, label, 11, MUTED)
footer(s, 2)

# ---- 3. Why now -----------------------------------------------------------
s = slide(); eyebrow(s, "Why now"); title(s, "Two curves just crossed")
bullets(s, [
    ("APIs got faster-moving", "old providers kept years of backward compatibility; AI APIs deprecate on ~90-day cycles. Breaking changes went from rare to constant."),
    ("AI got good enough to fix code", "code-generation models can now read a changelog and rewrite the affected code reliably — not possible 18 months ago."),
    ("The pain is compounding", "more of the software world is built on fast-moving AI APIs every quarter."),
], y=Inches(2.2), gap=16)
tf = textbox(s, Inches(0.7), Inches(5.7), Inches(11.9), Inches(1.0))
_set(tf.paragraphs[0], "Autonomously absorbing breaking API changes is newly feasible — and newly necessary.", 18, INK, bold=True, italic=True)
footer(s, 3)

# ---- 4. Solution ----------------------------------------------------------
s = slide(); eyebrow(s, "The solution"); title(s, "KeelCat turns a changelog into reviewed pull requests")
steps = [("Watch","provider changelogs + webhooks"),
         ("Understand","structured change-set (parser first, LLM only when needed)"),
         ("Find","every affected call site, across repos & languages"),
         ("Fix","deterministic transforms + LLM for the hard cases"),
         ("Verify","syntax + completeness check, sandbox run"),
         ("PR","one reviewable pull request per repo")]
for i,(h,b) in enumerate(steps):
    col = i % 3; row = i // 3
    x = Inches(0.7)+Inches(4.05)*col; y = Inches(2.25)+Inches(1.9)*row
    c = rect(s, x, y, Inches(3.8), Inches(1.65))
    tf = textbox(s, x+Inches(0.25), y+Inches(0.16), Inches(3.4), Inches(1.4))
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = f"{i+1}. "; r.font.size=Pt(16); r.font.bold=True; r.font.color.rgb=ACCENT; r.font.name=FONT
    r = p.add_run(); r.text = h; r.font.size=Pt(16); r.font.bold=True; r.font.color.rgb=INK; r.font.name=FONT
    p2 = tf.add_paragraph(); _set(p2, b, 12, MUTED)
tf = textbox(s, Inches(0.7), Inches(6.35), Inches(12), Inches(0.5))
_set(tf.paragraphs[0], "A human always reviews. KeelCat never merges on its own.", 15, ACCENT, bold=True)
footer(s, 4)

# ---- 5. Product / demo ----------------------------------------------------
s = slide(); eyebrow(s, "Product"); title(s, "Live today at keelcat.in")
bullets(s, [
    ("Connect GitHub, pick repos", "GitHub App or token; multi-account."),
    ("Register the APIs you depend on", "changelog URLs, polling + webhooks."),
    ("Watch it work", "live dashboard of repos changing, PRs opened, and their status."),
    ("Sign in with GitHub, Google, or as a guest", "isolated per-user workspaces."),
], y=Inches(2.1))
rect(s, Inches(0.7), Inches(5.5), Inches(11.9), Inches(1.2), fill=CARD)
tf = textbox(s, Inches(0.95), Inches(5.6), Inches(11.4), Inches(1.0), anchor=MSO_ANCHOR.MIDDLE)
_set(tf.paragraphs[0], "Built & shipped: multi-tenant SaaS, GitHub App + webhooks, hybrid fix engine, sandbox verification, security pass, Razorpay billing.", 14, MUTED)
footer(s, 5)

# ---- 6. Founder-market fit / origin --------------------------------------
s = slide(); eyebrow(s, "Why us"); title(s, "We lived this problem")
tf = textbox(s, Inches(0.7), Inches(2.1), Inches(11.9), Inches(2.2))
p = tf.paragraphs[0]
_set(p, "\u201cAs an intern, a missed changelog took our app down. I shipped 35 pull requests in 9 hours to fix every broken call site by hand.\u201d", 22, INK, bold=True, italic=True)
tf = textbox(s, Inches(0.7), Inches(4.1), Inches(11.9), Inches(1.0))
_set(tf.paragraphs[0], "That night is why KeelCat exists. It should have been one reviewed PR, opened automatically.", 16, MUTED)
bullets(s, [
    ("Deeply technical founders", "we build the full stack — detection, fix engine, and infra."),
    ("Shipping fast", "MVP live, deployed, and in the hands of real teams."),
], y=Inches(5.15))
footer(s, 6)

# ---- 7. Market ------------------------------------------------------------
s = slide(); eyebrow(s, "Market"); title(s, "A large, growing, structurally-tailwinded market")
for i,(num,label) in enumerate([("$8–10B","API management market, 2026"),
                                ("17–25%","annual growth"),
                                ("$240M+","ARR at 100k teams × ~$200/mo")]):
    x = Inches(0.7)+Inches(4.05)*i
    rect(s, x, Inches(2.2), Inches(3.8), Inches(1.9), fill=CARD)
    tf = textbox(s, x+Inches(0.2), Inches(2.45), Inches(3.4), Inches(1.5), anchor=MSO_ANCHOR.MIDDLE)
    _set(tf.paragraphs[0], num, 34, ACCENT, bold=True)
    p = tf.add_paragraph(); _set(p, label, 13, MUTED)
bullets(s, [
    ("Same buyer & budget", "as Dependabot, Snyk, and Sourcegraph."),
    ("Structural tailwind", "fast-moving AI APIs raise the frequency and cost of breaking changes every quarter."),
], y=Inches(4.5))
footer(s, 7)

# ---- 8. Competition -------------------------------------------------------
s = slide(); eyebrow(s, "Competition"); title(s, "The trigger, not the transform")
rows = [
    ("Version bots — Dependabot, Renovate, Snyk", "Bump versions, flag CVEs", "Don't read changelogs or rewrite your API code"),
    ("Code transforms — Sourcegraph, Moderne, Grit", "Large-scale migrations", "Engineer must author the recipe and trigger it"),
    ("KeelCat", "Detect a breaking change → fix → verify → PR, autonomously", "No human needed to trigger or author the migration"),
]
y = Inches(2.15)
for i,(a,b,c) in enumerate(rows):
    hl = (i == 2)
    r = rect(s, Inches(0.7), y, Inches(11.9), Inches(1.25), fill=(ACCENT if hl else CARD))
    tf = textbox(s, Inches(0.95), y+Inches(0.12), Inches(4.4), Inches(1.05), anchor=MSO_ANCHOR.MIDDLE)
    _set(tf.paragraphs[0], a, 15, (BG if hl else INK), bold=True)
    tf = textbox(s, Inches(5.5), y+Inches(0.12), Inches(3.3), Inches(1.05), anchor=MSO_ANCHOR.MIDDLE)
    _set(tf.paragraphs[0], b, 12, (BG if hl else MUTED))
    tf = textbox(s, Inches(9.0), y+Inches(0.12), Inches(3.4), Inches(1.05), anchor=MSO_ANCHOR.MIDDLE)
    _set(tf.paragraphs[0], c, 12, (BG if hl else MUTED))
    y += Inches(1.42)
tf = textbox(s, Inches(0.7), Inches(6.5), Inches(12), Inches(0.5))
_set(tf.paragraphs[0], "Scoping to external APIs makes autonomy tractable and accurate.", 15, ACCENT, bold=True)
footer(s, 8)

# ---- 9. Business model ----------------------------------------------------
s = slide(); eyebrow(s, "Business model"); title(s, "SaaS, bottom-up, expands with dependencies")
tiers = [("Free","Individuals & OSS","10 runs to try, then upgrade"),
         ("Team","$20–40 / dev / mo","unlimited runs, all repos"),
         ("Enterprise","Custom","internal APIs, policies, support")]
for i,(name,price,desc) in enumerate(tiers):
    x = Inches(0.7)+Inches(4.05)*i
    rect(s, x, Inches(2.2), Inches(3.8), Inches(2.4), fill=CARD, line=(ACCENT if i==1 else None))
    tf = textbox(s, x+Inches(0.3), Inches(2.45), Inches(3.2), Inches(2.0))
    _set(tf.paragraphs[0], name, 20, ACCENT, bold=True)
    p = tf.add_paragraph(); p.space_before=Pt(6); _set(p, price, 16, INK, bold=True)
    p = tf.add_paragraph(); p.space_before=Pt(6); _set(p, desc, 13, MUTED)
tf = textbox(s, Inches(0.7), Inches(5.0), Inches(11.9), Inches(1.0))
_set(tf.paragraphs[0], "Wedge: breaking-change auto-fix. Expansion: the system of record for every external dependency your code calls. Razorpay billing is live.", 15, MUTED)
footer(s, 9)

# ---- 10. Traction ---------------------------------------------------------
s = slide(); eyebrow(s, "Traction"); title(s, "Live, in the market, with first revenue")
for i,(num,label) in enumerate([("LIVE","product in production\nat keelcat.in"),
                                ("2 teams","actively testing"),
                                ("1st $","first paying user")]):
    x = Inches(0.7)+Inches(4.05)*i
    rect(s, x, Inches(2.2), Inches(3.8), Inches(1.9), fill=CARD)
    tf = textbox(s, x+Inches(0.2), Inches(2.42), Inches(3.4), Inches(1.5), anchor=MSO_ANCHOR.MIDDLE)
    _set(tf.paragraphs[0], num, 30, ACCENT, bold=True)
    p = tf.add_paragraph(); _set(p, label, 12, MUTED)
bullets(s, [
    ("Shipped end-to-end", "multi-tenant auth, GitHub/Google/guest login, billing, deployment."),
    ("North-star metric", "PR merge rate — are teams actually merging what we open?"),
], y=Inches(4.5))
footer(s, 10)

# ---- 11. Team -------------------------------------------------------------
s = slide(); eyebrow(s, "Team"); title(s, "Builders who ship")
people = [("Vinod","Builds the full stack — detection, fix engine, infra. Startups since 17; hardware, trading systems, and autonomous systems background."),
          ("Srushti","Co-founder; went full-time on KeelCat. Engineering classmates since 2022.")]
for i,(name,bio) in enumerate(people):
    y = Inches(2.2)+Inches(1.75)*i
    rect(s, Inches(0.7), y, Inches(11.9), Inches(1.5), fill=CARD)
    tf = textbox(s, Inches(1.0), y+Inches(0.18), Inches(11.3), Inches(1.15), anchor=MSO_ANCHOR.MIDDLE)
    _set(tf.paragraphs[0], name, 18, ACCENT, bold=True)
    p = tf.add_paragraph(); _set(p, bio, 14, MUTED)
footer(s, 11)

# ---- 12. Ask / vision -----------------------------------------------------
s = slide()
rect(s, Inches(0.7), Inches(2.2), Inches(0.14), Inches(3.0), fill=ACCENT)
tf = textbox(s, Inches(1.05), Inches(2.15), Inches(11), Inches(1.0))
_set(tf.paragraphs[0], "Join Techstars with us", 34, INK, bold=True)
bullets(s, [
    ("Now", "grow the design-partner base and drive PR merge rate up."),
    ("Next", "broader language coverage, pre-wired provider registry, run each repo's own tests."),
    ("Vision", "the system that keeps the world's code working as the APIs beneath it change."),
], x=Inches(1.05), y=Inches(3.25), gap=14)
tf = textbox(s, Inches(1.05), Inches(6.2), Inches(11), Inches(0.6))
p = tf.paragraphs[0]
r = p.add_run(); r.text = "keelcat.in"; r.font.size=Pt(16); r.font.bold=True; r.font.color.rgb=ACCENT; r.font.name=FONT
r = p.add_run(); r.text = "     ·     Vinod & Srushti     ·     Your APIs change. Your code keeps up."; r.font.size=Pt(14); r.font.color.rgb=MUTED; r.font.name=FONT

out = __file__.replace("build_pitch_deck.py", "KeelCat-Pitch-Deck.pptx")
prs.save(out)
print("Saved:", out, "-", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
